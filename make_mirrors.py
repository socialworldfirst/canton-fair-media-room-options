# -*- coding: utf-8 -*-
import json,math
exec(open('venues.py',encoding='utf-8').read())
IMG=json.load(open('images.json',encoding='utf-8')); RT=json.load(open('routes.json',encoding='utf-8'))
km=lambda ll: round(math.hypot((ll[1]-REF[1])*102400,(ll[0]-REF[0])*111320)/1000,1)
mins=lambda s: max(1,int(round(s/60.0)))
GDOC="https://docs.google.com/document/d/1ygD-NcaK9cDrzkl9Cp2NQSzEUdPw-Pb018rzq4zat5A/edit"
order=sorted(V,key=lambda v:(km(v['ll']),v['rank']))
def tstr(vid):
    r=RT[vid]
    if not r.get('walk_s'): return "already in the building","0 min","0 min"
    return ("%d min walk (%.1f km)"%(mins(r['walk_s']),r['walk_m']/1000.0),
            "%d min drive (%.1f km)"%(mins(r['drive_s']),r['drive_m']/1000.0),
            "%d min in fair traffic"%mins(r['drive_s']*2.2))
def gmaps(ll): return "https://www.google.com/maps/search/?api=1&query=%.5f,%.5f"%ll
def gdir(ll,walk): return "https://www.google.com/maps/dir/?api=1&origin=23.10344,113.35412&destination=%.5f,%.5f&travelmode=%s"%(ll[0],ll[1],'walking' if walk else 'driving')

md=["# 广交会 2026 · 媒体间选址 20 个方案","",
"**Canton Fair 2026 (140th session) · media room location options**","",
"Twenty shortlisted rooms, halls, decks and rooftops for interviews and filmed conversations around the China Import and Export Fair Complex in Pazhou, Guangzhou.","",
"Fair dates: Phase 1 15–19 Oct · Phase 2 23–27 Oct · Phase 3 31 Oct – 4 Nov 2026.","",
"Travel times and route distances come from OpenStreetMap road routing, measured from the complex on Yuejiang Zhong Road. Fair traffic is a 2.2x allowance on the free-flow drive. Nothing here is booked.","",
"> **配图 + 路线图版 / with photos and route maps:** %s"%GDOC,
"> 每个场地 3 张实景图、从广交会出发的路线地图、Google Maps 链接都在这份 Google Doc 里。","",
"| # | Venue 场地 | Zone | Walk | Drive | Fair traffic | Google Maps |","|---|---|---|---|---|---|---|"]
for i,v in enumerate(order,1):
    w,d,t=tstr(v['id']); r=RT[v['id']]
    md.append("| %02d | %s | %s | %s | %s | %s | [pin](%s) |"%(i,v['name'],v['zone'],w,d,t,gmaps(tuple(v['ll']))))
md.append("")
for i,v in enumerate(order,1):
    w,d,t=tstr(v['id']); r=RT[v['id']]; walk=r.get('mode')=='walk'
    md += ["## %02d. %s"%(i,v['name']), "",
           "**%s** · %s · %s km straight line"%(v['cn'],v['zone'],km(v['ll'])), "",
           "- **地址 Address:** %s"%v['addr'],
           "- **步行 Walk:** %s"%w,
           "- **驾车 Drive:** %s，展会期间按 %s 预留"%(d,t),
           "- **画面 Backdrop:** %s"%v['backdrop'],
           "- **坐标 Coordinates:** %.5f, %.5f"%(v['ll'][0],v['ll'][1]),
           "- **Google Maps:** [定位 pin](%s) · [从广交会出发 directions](%s)"%(gmaps(tuple(v['ll'])),gdir(v['ll'],walk)),
           "- **标签 Tags:** %s"%" / ".join(v['tags']), "",
           "**Why it works.** %s"%v['why'], "",
           "**Watch out.** %s"%v['watch'], ""]
md += ["---","","**Shortlists**","",
 "- 如果主角是广交会 If the story is the fair: Official Press Centre (Area A) · Area D conference rooms · Nan Fung exhibition halls",
 "- 如果主角是天际线 If the story is the skyline: Guangzhou International Media Port · Park Hyatt Roof Bar · Haixin Bridge riverfront",
 "- 最稳妥、最快点头 Lowest risk, fastest yes: Shangri-La Hotel Guangzhou · InterContinental Guangzhou Exhibition Center · Langham Place Guangzhou","",
 "**待确认 Still to confirm:** fair-week availability at every hotel · filming permits for public riverside sites · GRT studio access and a media partner.","",
 "*Route maps are OpenStreetMap. Photography is Creative Commons and public domain, from Wikimedia Commons and Openverse, credited in the Google Doc. Venue details come from public sources and are not confirmed bookings.*"]
open('yuque-广交会2026-媒体间选址.md','w',encoding='utf-8').write("\n".join(md).replace('$','\\$'))

from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
doc=Document()
for s in doc.sections:
    s.left_margin=s.right_margin=Inches(0.7); s.top_margin=s.bottom_margin=Inches(0.6)
doc.styles['Normal'].font.name='Poppins'; doc.styles['Normal'].font.size=Pt(10)
doc.add_heading('Canton Fair 2026 · 20 media room location options',0)
doc.add_paragraph('广交会 2026（第140届）· 媒体间选址 20 个方案')
doc.add_paragraph('Twenty shortlisted rooms, halls, decks and rooftops for interviews and filmed conversations around the China Import and Export Fair Complex in Pazhou, Guangzhou. Each option carries a route map from the complex, real walking and driving times from OpenStreetMap routing, and Google Maps links. Fair traffic is a 2.2x allowance on the free-flow drive. Nothing here is booked.')
doc.add_paragraph('Fair dates: Phase 1 15–19 Oct · Phase 2 23–27 Oct · Phase 3 31 Oct – 4 Nov 2026.')
doc.add_heading('At a glance',1)
t=doc.add_table(rows=1,cols=5); t.style='Light Grid Accent 1'
for c,h in zip(t.rows[0].cells,['#','Venue','Zone','Walk','Drive in fair traffic']): c.text=h
for i,v in enumerate(order,1):
    w,d,tt=tstr(v['id']); row=t.add_row().cells
    row[0].text='%02d'%i; row[1].text=v['name']; row[2].text=v['zone']; row[3].text=w; row[4].text=tt
for i,v in enumerate(order,1):
    doc.add_page_break()
    doc.add_heading('%02d. %s'%(i,v['name']),1)
    p=doc.add_paragraph(); r=p.add_run('%s · %s · %s km straight line'%(v['cn'],v['zone'],km(v['ll'])))
    r.font.color.rgb=RGBColor(0xFF,0x00,0x51); r.font.size=Pt(9.5)
    shots=IMG.get(v['id'],[])[:3]
    if shots:
        tb=doc.add_table(rows=1,cols=len(shots))
        for c,s in zip(tb.rows[0].cells,shots):
            cp=c.paragraphs[0]; cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            try: cp.add_run().add_picture(s['file'],width=Inches(6.9/len(shots)))
            except Exception: pass
            cap=c.add_paragraph(); cr=cap.add_run('%s — %s (%s)'%((s.get('title') or '')[:48],(s.get('credit') or '')[:26],(s.get('lic') or '').strip()))
            cr.font.size=Pt(6); cr.font.color.rgb=RGBColor(0x8A,0x8A,0x8A)
    mp=doc.add_paragraph(); mp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    try: mp.add_run().add_picture('maps/%s.jpg'%v['id'],width=Inches(6.9))
    except Exception: pass
    mc=doc.add_paragraph(); mcr=mc.add_run('Route from the Canton Fair Complex. Map © OpenStreetMap contributors.')
    mcr.font.size=Pt(6.5); mcr.font.color.rgb=RGBColor(0x8A,0x8A,0x8A); mc.alignment=WD_ALIGN_PARAGRAPH.CENTER
    w,d,tt=tstr(v['id']); walk=RT[v['id']].get('mode')=='walk'
    for lab,val in [('Address',v['addr']),('Walk',w),('Drive',"%s, allow %s during the fair"%(d,tt)),
                    ('Backdrop',v['backdrop']),('Coordinates','%.5f, %.5f'%(v['ll'][0],v['ll'][1])),
                    ('Google Maps','%s   |   directions: %s'%(gmaps(tuple(v['ll'])),gdir(v['ll'],walk))),
                    ('Tags',' / '.join(v['tags']))]:
        pp=doc.add_paragraph(); pp.add_run(lab+': ').bold=True; rr=pp.add_run(val)
        if lab=='Google Maps': rr.font.size=Pt(7.5)
    pp=doc.add_paragraph(); pp.add_run('Why it works. ').bold=True; pp.add_run(v['why'])
    pp=doc.add_paragraph(); pp.add_run('Watch out. ').bold=True; pp.add_run(v['watch'])
doc.add_page_break(); doc.add_heading('Shortlists',1)
for ttl,items in [('If the story is the fair',['Official Press Centre (Area A)','Area D conference rooms','Nan Fung exhibition halls']),
                ('If the story is the skyline',['Guangzhou International Media Port','Park Hyatt Roof Bar','Haixin Bridge riverfront']),
                ('Lowest risk, fastest yes',['Shangri-La Hotel, Guangzhou','InterContinental Guangzhou Exhibition Center','Langham Place, Guangzhou']),
                ('Still to confirm',['Fair-week availability at every hotel','Filming permits for public riverside sites','GRT studio access and a media partner'])]:
    doc.add_heading(ttl,2)
    for it in items: doc.add_paragraph(it,style='List Bullet')
doc.add_paragraph('Route maps © OpenStreetMap contributors. Photography is Creative Commons and public domain, from Wikimedia Commons and Openverse, credited under each image. Venue details come from public sources and are not confirmed bookings.')
doc.save('Canton-Fair-2026-media-room-options.docx')
print('done')
